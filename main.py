import cv2
import pytesseract
import pandas as pd
import numpy as np
import re
import os
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ruta a Tesseract en Windows
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

def enderezar_imagen(gris):
    """Detecta y corrige inclinaciones leves en la imagen."""
    _, thresh = cv2.threshold(gris, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    coords = np.column_stack(np.where(thresh > 0))
    if len(coords) > 0:
        angle = cv2.minAreaRect(coords)[-1]
        if angle < -45:
            angle = -(90 + angle)
        elif angle > 45:
            angle = 90 - angle
        else:
            angle = -angle
        
        if 0.5 < abs(angle) < 15:
            (h, w) = gris.shape[:2]
            centro = (w // 2, h // 2)
            M = cv2.getRotationMatrix2D(centro, angle, 1.0)
            gris = cv2.warpAffine(gris, M, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)
    return gris

def guardar_excel_estilizado(df: pd.DataFrame, ruta_salida: str):
    """Guarda el DataFrame en un archivo Excel con diseño profesional."""
    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Lockers", startrow=3)
        ws = writer.sheets["Lockers"]
        
        ws.views.sheetView[0].showGridLines = True

        # 1. Título principal
        ws.merge_cells("A1:B1")
        ws["A1"] = "CONTROL Y ASIGNACIÓN DE LOCKERS"
        ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        
        # Subtítulo con fecha y total
        fecha_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        ws.merge_cells("A2:B2")
        ws["A2"] = f"Generado: {fecha_str}  |  Total Asignados: {len(df)}"
        ws["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="595959")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        # 2. Encabezados de columna
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        
        for col_num in range(1, len(df.columns) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 25

        # 3. Bordes y estilo cebra
        thin_border = Border(
            left=Side(style='thin', color="D9D9D9"),
            right=Side(style='thin', color="D9D9D9"),
            top=Side(style='thin', color="D9D9D9"),
            bottom=Side(style='thin', color="D9D9D9")
        )
        zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        start_row = 5
        for row_idx, row in enumerate(ws.iter_rows(min_row=start_row, max_row=start_row + len(df) - 1, min_col=1, max_col=len(df.columns))):
            ws.row_dimensions[row[0].row].height = 20
            es_par = (row_idx % 2 == 1)
            for col_idx, cell in enumerate(row):
                cell.border = thin_border
                cell.font = Font(name="Segoe UI", size=10)
                cell.fill = zebra_fill if es_par else white_fill
                
                if col_idx == 0:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # 4. Autoajustar ancho de columnas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row >= 4 and cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 6, 16)

def procesar_hoja_lockers(ruta_imagen, salida_excel="lockers_limpio.xlsx"):
    if not os.path.exists(ruta_imagen):
        print(f"No se encontró el archivo: {ruta_imagen}")
        return

    # 1. Cargar imagen completa
    img = cv2.imread(ruta_imagen)
    gris = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)

    # 2. Enderezar imagen si está ligeramente inclinada
    gris = enderezar_imagen(gris)

    # 3. Reescalar a x2 de resolución para máxima definición de caracteres
    gris_2x = cv2.resize(gris, None, fx=2.0, fy=2.0, interpolation=cv2.INTER_CUBIC)

    # 4. Umbralización Otsu de alta fidelidad
    _, thresh = cv2.threshold(gris_2x, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)

    # 5. OCR por líneas completas (PSM 6: bloque uniforme de texto)
    config = r'--oem 3 --psm 6 -l spa'
    texto_crudo = pytesseract.image_to_string(thresh, config=config)

    registros = []
    ultimo_num = 0

    for linea in texto_crudo.splitlines():
        linea = linea.strip()
        if not linea or "LOCKER" in linea.upper():
            continue

        # Limpiar símbolos que suelen confundirse con las líneas de la cuadrícula
        linea_limpia = re.sub(r'[\[\]\|\_\—\-\~\{\}\:\;\¿\¡\°\º\(\)\<\>\"\=\'\?\*\+\#\$\%]', ' ', linea)
        linea_limpia = " ".join(linea_limpia.split())

        if len(linea_limpia) < 4:
            continue

        # Buscar todos los posibles números en la línea
        nums = re.findall(r'\b\d{1,3}\b', linea_limpia)
        num_elegido = None
        
        for n_str in nums:
            n = int(n_str)
            # Manejar errores de OCR en números pegados como '117' para el 11
            candidatos = [n]
            if n > 50:
                candidatos.extend([n // 10, n % 100])
                if len(n_str) >= 2:
                    candidatos.append(int(n_str[:2]))
            
            for cand in candidatos:
                if ultimo_num < cand <= ultimo_num + 3:
                    num_elegido = cand
                    break
            if num_elegido is not None:
                break

        if num_elegido is None:
            num_elegido = ultimo_num + 1

        # Extraer el nombre (donde comienzan las palabras)
        m_letras = re.search(r'[A-ZÁÉÍÓÚÑa-záéíóúñ]{3,}.*', linea_limpia)
        if not m_letras:
            continue
        
        nombre = m_letras.group(0)
        
        # Limpieza de prefijos como 'ENTER 2' o letras sueltas
        nombre = re.sub(r'^(ENTER|\d+)\s+', '', nombre, flags=re.IGNORECASE).strip()
        nombre = re.sub(r'^[a-zA-Z]{1,2}\s+', '', nombre).strip()

        # Limpieza de anotaciones manuscritas al final
        nombre = re.sub(r'\b(NO|OK|O|AA|TN|WD|TS|OS)\b$', '', nombre, flags=re.IGNORECASE).strip()
        nombre = re.sub(r'[\d\W]+$', '', nombre).strip()
        nombre = re.sub(r'\s+[A-Z0-9]{1,2}$', '', nombre).strip()
        nombre = " ".join(nombre.split()).upper()

        palabras = nombre.split()
        if len(nombre) >= 4 and len(palabras) >= 2 and any(c.isalpha() for c in nombre):
            ultimo_num = num_elegido
            registros.append({
                "N° LOCKER": num_elegido,
                "NOMBRE": nombre
            })

    df = pd.DataFrame(registros)
    if not df.empty:
        df = df.drop_duplicates(subset=["N° LOCKER"]).sort_values(by="N° LOCKER")
        guardar_excel_estilizado(df, salida_excel)
        print(f"¡Éxito! Se extrajeron {len(df)} lockers con diseño profesional en '{salida_excel}'.")
    else:
        print("No se encontraron registros válidos. Revisa la calidad de la foto.")

if __name__ == "__main__":
    procesar_hoja_lockers("Hoja.jpeg")